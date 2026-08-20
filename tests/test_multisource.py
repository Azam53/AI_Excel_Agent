from io import BytesIO
from openpyxl import load_workbook

from connectors.base import ConnectorError
from connectors.fred import FredConnector
from services.csv_generator import generate_csv
from services.data_merger import merge_results
from services.excel_generator import generate_multi_workbook
from services.parser import merge_context, parse_request
from services.query_planner import build_plan, execute_plan
from services.source_router import SourceRouter

def test_detects_multiple_metrics_and_sources():
    parsed = parse_request("Compare India's GDP, inflation and Brent crude oil prices from 2015 to 2025", 2026)
    assert parsed["metrics"] == ["gdp", "inflation", "brent_crude"]
    plan = build_plan(parsed)
    assert [task["source"] for task in plan["tasks"]] == ["worldbank", "worldbank", "fred"]

def test_planner_expands_world_bank_by_country():
    context = parse_request("Compare India and UAE GDP and inflation from 2018 to 2025", 2026)
    plan = build_plan(context)
    assert len(plan["tasks"]) == 4
    assert {task["country"]["code"] for task in plan["tasks"]} == {"IND", "ARE"}

def test_router_uses_registry():
    router = SourceRouter()
    assert router.route("gdp")[0] == "worldbank"
    assert router.route("brent_crude")[0] == "fred"

def test_merge_keeps_missing_values_blank():
    results = [
        {"source":"World Bank","metric":"GDP","metric_key":"gdp","country":"India","unit":"USD","frequency":"annual","records":[{"period":2020,"value":1},{"period":2021,"value":None}],"metadata":{}},
        {"source":"FRED","metric":"Brent Crude Oil","metric_key":"brent_crude","country":None,"unit":"USD/barrel","frequency":"annual","records":[{"period":2020,"value":40},{"period":2021,"value":70}],"metadata":{"aggregation":"Annual average"}},
    ]
    merged = merge_results(results, 2020, 2021)
    assert merged["rows"][1]["India GDP"] is None
    assert merged["rows"][1]["Brent Crude Oil"] == 70

def test_follow_up_context_adds_country_metric_and_after_year():
    context = parse_request("Compare India and UAE GDP from 2015 to 2025", 2026)
    context = merge_context(context, parse_request("Add Saudi Arabia", 2026, require_complete=False), 2026)
    context = merge_context(context, parse_request("Add inflation too", 2026, require_complete=False), 2026)
    context = merge_context(context, parse_request("Only show data after 2020", 2026, require_complete=False), 2026)
    assert [country["code"] for country in context["countries"]] == ["IND", "ARE", "SAU"]
    assert context["metrics"] == ["gdp", "inflation"]
    assert (context["start_year"], context["end_year"]) == (2021, 2025)

def test_csv_and_multi_metric_excel():
    context = parse_request("Compare India GDP and inflation from 2020 to 2021", 2026)
    results = [
        {"source":"World Bank","metric":"GDP","metric_key":"gdp","country":"India","unit":"USD","frequency":"annual","records":[{"period":2020,"value":1},{"period":2021,"value":2}],"metadata":{"indicator":"GDP","source_url":"https://example.test/gdp"}},
        {"source":"World Bank","metric":"Inflation","metric_key":"inflation","country":"India","unit":"%","frequency":"annual","records":[{"period":2020,"value":5},{"period":2021,"value":6}],"metadata":{"indicator":"CPI","source_url":"https://example.test/cpi"}},
    ]
    merged = merge_results(results, 2020, 2021)
    assert generate_csv(merged).getvalue().decode("utf-8-sig").startswith("Year,India GDP,India Inflation")
    workbook = load_workbook(BytesIO(generate_multi_workbook(context, results, merged).getvalue()))
    assert workbook.sheetnames == ["Dashboard", "Combined Data", "Sources"]
    assert len(workbook["Dashboard"]._charts) == 2

def test_missing_fred_is_partial_not_crash():
    class Good:
        name = "World Bank"
        def fetch(self, task): return {"source":"World Bank","metric":"GDP","metric_key":"gdp","country":"India","unit":"USD","frequency":"annual","records":[{"period":2020,"value":1}],"metadata":{}}
        def status(self): return {"name":self.name,"available":True,"message":"Connected"}
    class Missing:
        name = "FRED"
        def fetch(self, task): raise ConnectorError("FRED unavailable because no API key is configured.")
        def status(self): return {"name":self.name,"available":False,"message":"API key not configured"}
    router = SourceRouter({"worldbank":Good(), "fred":Missing()})
    plan = {"tasks":[{"metric":"gdp","source":"worldbank"},{"metric":"brent_crude","source":"fred"}]}
    results, errors = execute_plan(plan, router)
    assert len(results) == 1 and len(errors) == 1
    assert "no API key" in errors[0]["error"]

def test_fred_without_key_reports_unavailable():
    connector = FredConnector(api_key="")
    assert connector.status()["available"] is False
    try: connector.fetch({"metric":"brent_crude","start_year":2020,"end_year":2021})
    except ConnectorError as exc: assert "no API key" in str(exc)
    else: raise AssertionError("Expected unavailable FRED connector")

def test_fred_with_key_aggregates_to_annual_average():
    class Response:
        def raise_for_status(self): pass
        def json(self): return {"observations":[{"date":"2020-01-01","value":"40"},{"date":"2020-07-01","value":"60"},{"date":"2021-01-01","value":"."}]}
    class Session:
        @staticmethod
        def get(*args, **kwargs): return Response()
    result = FredConnector(api_key="demo", session=Session).fetch({"metric":"brent_crude","start_year":2020,"end_year":2021})
    assert result["records"] == [{"period":2020,"value":50.0},{"period":2021,"value":None}]
    assert result["metadata"]["aggregation"] == "Annual average of available observations"

def test_fred_only_plan_does_not_require_country():
    parsed = parse_request("Show Brent crude oil prices from 2020 to 2024", 2026)
    assert parsed["countries"] == []
    assert build_plan(parsed)["tasks"][0]["source"] == "fred"

def test_chat_guides_user_when_metric_is_missing():
    from app import app
    client = app.test_client()
    response = client.post("/api/chat", json={"message":"Tell me something about India"})
    payload = response.get_json()
    assert response.status_code == 400
    assert "measurable topic" in payload["error"]
    assert len(payload["suggestions"]) >= 2
    assert payload["can_retry"] is True

def test_chat_guides_user_when_country_is_missing():
    from app import app
    client = app.test_client()
    response = client.post("/api/chat", json={"message":"Show GDP from 2020 to 2024"})
    payload = response.get_json()
    assert response.status_code == 400
    assert "Add a country" in payload["error"]
    assert payload["suggestions"]
