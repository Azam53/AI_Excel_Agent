"""Export source listings, never statistical observations."""
import csv
from io import BytesIO, StringIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def safe_cell(value):
    value = str(value or '')
    value = ''.join(c for c in value if ord(c) >= 32 or c in '\n\t')
    return "'" + value if value.lstrip().startswith(('=', '+', '-', '@')) else value


def generate_search_exports(question, research):
    headers = ['Title', 'Source', 'Domain', 'Snippet', 'Published Date', 'URL']
    keys = ['title', 'source', 'domain', 'snippet', 'published_date', 'url']
    rows = [[safe_cell(result.get(key)) for key in keys] for result in research['results']]
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Search Results'
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in zip('ABCDEF', [45, 25, 25, 90, 18, 65]):
        sheet.column_dimensions[column].width = width
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for row in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 100
        sheet.cell(row, 6).hyperlink = research['results'][row - 2]['url']
        sheet.cell(row, 6).style = 'Hyperlink'
    info = wb.create_sheet('Search Info')
    info.append(['Field', 'Value'])
    for key, value in [('Original Question', question), ('Search Query', research['query']),
                       ('Generated Date', research['generated_at']), ('Number of Results', len(rows)),
                       ('Search Provider', research['provider']),
                       ('Notice', 'Search snippets are source listings, not verified facts or a statistical dataset.')]:
        info.append([key, safe_cell(value)])
    info.column_dimensions['A'].width = 25
    info.column_dimensions['B'].width = 100
    for ws in wb:
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='173D55')
    excel = BytesIO()
    wb.save(excel)
    excel.seek(0)
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows(rows)
    return excel, BytesIO(output.getvalue().encode('utf-8-sig'))
