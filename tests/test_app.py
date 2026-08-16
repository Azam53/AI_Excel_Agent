from io import BytesIO
from openpyxl import load_workbook
from services.parser import parse_prompt
from services.worldbank import fetch_country_data
from services.excel_generator import generate_workbook

def test_country_indicator_and_year_parsing():
    result=parse_prompt("Show India's inflation from 2015 to 2025",2026)
    assert result.countries[0]["code"]=="IND" and result.indicator_name=="Inflation"
    assert (result.start_year,result.end_year)==(2015,2025)

def test_gdp_growth_precedes_gdp():
    assert parse_prompt("Compare GDP growth of India, China and USA from 2015 to 2025",2026).indicator_code=="NY.GDP.MKTP.KD.ZG"

def test_last_years():
    result=parse_prompt("Show life expectancy in India for the last 15 years",2026)
    assert (result.start_year,result.end_year)==(2012,2026)

class Response:
    def raise_for_status(self): pass
    def json(self): return [{},[{"date":"2024","value":123.5},{"date":"2023","value":None}]]
class Session:
    last_kwargs = None
    @staticmethod
    def get(*args,**kwargs):
        Session.last_kwargs = kwargs
        return Response()

def test_world_bank_response_parsing():
    values,url=fetch_country_data("IND","NY.GDP.MKTP.CD",2023,2024,Session)
    assert values=={2024:123.5,2023:None} and "country/IND" in url
    assert Session.last_kwargs["timeout"] == (10, 60)
    assert Session.last_kwargs["headers"]["Accept"] == "application/json"

def test_workbook_creation():
    parsed=parse_prompt("Compare India and UAE GDP from 2020 to 2022",2026)
    data={"India":{2020:1.0,2021:2.0,2022:3.0},"United Arab Emirates":{2020:4.0,2021:5.0,2022:6.0}}
    wb=load_workbook(BytesIO(generate_workbook(parsed,data,["https://example.test"]).getvalue()))
    assert wb.sheetnames==["Dashboard","Data","Sources"] and len(wb["Dashboard"]._charts)==1

def test_single_country_chart_and_unavailable_year_note():
    parsed=parse_prompt("Show life expectancy in India for the last 15 years",2026)
    data={"India":{year:(70.0 + (year-2012)*0.1 if year <= 2024 else None) for year in range(2012,2027)}}
    wb=load_workbook(BytesIO(generate_workbook(parsed,data,["https://example.test"]).getvalue()))
    dashboard=wb["Dashboard"]
    assert len(dashboard._charts)==1
    assert dashboard._charts[0].title.tx.rich.p[0].r[0].t == "Life Expectancy — India — 2012 to 2026"
    assert wb["Data"]["B15"].value is None
    assert any(cell.value == "Latest World Bank data available: 2024" for row in dashboard.iter_rows() for cell in row)
