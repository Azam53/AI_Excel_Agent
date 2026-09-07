from io import BytesIO
from unittest.mock import Mock
import pytest
import requests
from openpyxl import load_workbook
import app as application
from connectors.web_search import WebSearchConnector, WebSearchError, normalize_results, safe_url, source_authority
from services.request_interpreter import interpret_request
from services.search_query_builder import build_search_query


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setenv('WEB_SEARCH_PROVIDER', 'tavily')
    monkeypatch.setenv('WEB_SEARCH_API_KEY', 'test-secret')
    application.app.config['TESTING'] = True
    return application.app.test_client()


@pytest.fixture
def search(monkeypatch):
    result = normalize_results('India AI market size', [{'title': '=India AI market', 'url': 'https://india.gov.in/report', 'content': '<script>alert(1)</script> AI market research', 'published_date': '2026-01-01'}])
    mock = Mock(side_effect=lambda query, **kw: {'query': query, 'provider': 'tavily', 'results': result})
    monkeypatch.setattr(WebSearchConnector, 'search', mock)
    return mock


def mock_structured(monkeypatch):
    def execute(plan):
        from services.registry import METRICS
        return [{'source': 'World Bank' if task['source'] == 'worldbank' else 'FRED',
                 'metric_key': task['metric'], 'metric': METRICS[task['metric']]['label'],
                 'country': task.get('country', {}).get('name', 'Global'),
                 'unit': 'USD', 'frequency': 'annual', 'source_url': 'https://data.worldbank.org/',
                 'records': [{'period': 2024, 'value': 100}], 'metadata': {}}
                for task in plan['tasks']], []
    monkeypatch.setattr(application, 'execute_plan', execute)


@pytest.mark.parametrize('prompt', [
    'Compare India and UAE GDP from 2015 to 2025',
    'Show life expectancy in India for the last 15 years',
    "Compare India's GDP with Brent crude oil prices from 2015 to 2025",
])
def test_structured_never_searches(client, search, monkeypatch, prompt):
    mock_structured(monkeypatch)
    response = client.post('/api/chat', json={'message': prompt})
    assert response.status_code == 200, response.get_json()
    data = response.get_json()
    assert data['mode'] == 'structured' and data['excel_available']
    assert client.get('/api/download/' + data['download_token'] + '/excel').status_code == 200
    search.assert_not_called()
    plan = client.post('/api/chat/plan', json={'message': prompt}).get_json()
    assert not any('Searching' in step for step in plan['steps'])


def test_unsupported_search_and_export(client, search):
    data = client.post('/api/chat', json={'message': "What is India's AI market size?"}).get_json()
    assert data['mode'] == 'web' and not data['excel_available']
    assert data['web']['query'] == 'India AI market size latest'
    token = data['web']['download_token']
    excel = client.get('/api/download/' + token + '/excel')
    wb = load_workbook(BytesIO(excel.data))
    assert wb.sheetnames == ['Search Results', 'Search Info']
    assert wb['Search Results']['A2'].data_type != 'f'
    assert wb['Search Results']['F2'].hyperlink.target == 'https://india.gov.in/report'
    csv = client.get('/api/download/' + token + '/csv')
    assert b'Title,Source,Domain,Snippet,Published Date,URL' in csv.data
    assert b"'=India" in csv.data


def test_unavailable(client, monkeypatch):
    monkeypatch.delenv('WEB_SEARCH_API_KEY')
    data = client.post('/api/chat', json={'message': 'India AI market size'}).get_json()
    assert data['web']['status'] == 'not_configured'
    assert 'not currently configured' in data['message']


@pytest.mark.parametrize('outcome', ['error', 'empty'])
def test_search_failures(client, search, outcome):
    if outcome == 'error':
        search.side_effect = WebSearchError("I couldn't complete the web search right now. Please try again.")
    else:
        search.side_effect = lambda query: {'query': query, 'provider': 'tavily', 'results': []}
    response = client.post('/api/chat', json={'message': 'India AI market size'})
    assert response.status_code == 200
    assert response.get_json()['web']['status'] == outcome
    assert 'download_token' not in response.get_json()['web']


def test_mixed_and_follow_up_reset(client, search, monkeypatch):
    mock_structured(monkeypatch)
    data = client.post('/api/chat', json={'message': "Show India's GDP from 2018 to 2025 and find information about India's AI industry growth."}).get_json()
    assert data['mode'] == 'mixed' and data['excel_available'] and data['web']['results']
    assert data['context']['metrics'] == ['GDP']
    client.post('/api/chat', json={'message': "What is India's AI market size?"})
    data = client.post('/api/chat', json={'message': 'What about UAE?'}).get_json()
    assert data['web']['query'] == 'UAE AI market size latest'
    client.post('/api/chat/reset')
    with client.session_transaction() as state:
        assert 'web_context' not in state and 'conversation' not in state


def test_structured_failure_does_not_search(client, search, monkeypatch):
    monkeypatch.setattr(application, 'execute_plan', lambda plan: ([], [{'source':'World Bank', 'metric':'gdp', 'error':"I couldn't reach World Bank right now."}]))
    response = client.post('/api/chat', json={'message': 'Show India GDP'})
    assert response.status_code == 502
    search.assert_not_called()


@pytest.mark.parametrize('url', ['javascript:alert(1)', 'data:text/html,a', '//example.com', 'https://user:pass@example.com', 'http://127.0.0.1', 'https://example.com\\@evil.com', 'https://example.com\n', 'https://localhost', 'https://example.com:bad'])
def test_unsafe_urls(url):
    assert not safe_url(url)


def test_normalization_and_ranking():
    result = normalize_results('AI market', [
        {'url': 'javascript:bad', 'title': 'AI'},
        {'url': 'https://unknown.org/a', 'title': 'AI market'},
        {'url': 'https://imf.org/a', 'title': 'AI market', 'published_date': '2026-01-01'},
        {'url': 'https://imf.org/a', 'title': 'duplicate'}, None])
    assert len(result) == 2 and result[0]['domain'] == 'imf.org'
    assert result[0]['published_date'] == '2026-01-01'
    assert source_authority('unknown.org')[0] == 0
    assert source_authority('imf.org.evil.com')[0] == 0


def test_provider_adapter_and_errors():
    transport = Mock()
    transport.post.return_value.json.return_value = {'results': [{'url': 'https://imf.org/a', 'title': 'AI'}]}
    connector = WebSearchConnector('tavily', 'secret', transport)
    assert len(connector.search('AI market')['results']) == 1
    assert transport.post.call_args.kwargs['timeout'] == (5, 25)
    assert transport.post.call_args.kwargs['json']['include_answer'] is False
    transport.post.side_effect = requests.Timeout('secret')
    with pytest.raises(WebSearchError) as exc:
        connector.search('AI market')
    assert 'secret' not in str(exc.value)


def test_status(client):
    statuses = client.get('/api/sources').get_json()['sources']
    assert len(statuses) == 4
    assert statuses[-1]['type'] == 'search' and statuses[-1]['status'] == 'connected'
    assert statuses[0]['type'] == 'structured'


@pytest.mark.parametrize('payload', [[], 'text', {'message': 123}, {'message': ''}, {'message': 'x' * 501}, {'message': '!!!'}])
def test_invalid_input(client, search, payload):
    assert client.post('/api/chat', json=payload).status_code == 400
    search.assert_not_called()


def test_query_builder_examples():
    assert build_search_query('How much renewable energy does UAE produce?')[0] == 'UAE renewable energy production statistics latest'
    assert build_search_query("What is India's ecommerce market expected to reach by 2030?")[0] == 'India ecommerce market size forecast 2030'


def test_keyless_adapter_omits_credentials():
    transport = Mock()
    transport.post.return_value.json.return_value = {'results': []}
    connector = WebSearchConnector('tavily-keyless', 'unused-secret', transport)
    assert connector.is_available()
    assert 'rate limited' in connector.status()['message']
    connector.search('vibe coding latest')
    assert transport.post.call_args.kwargs['headers'] == {'X-Tavily-Access-Mode': 'keyless'}
    assert not WebSearchConnector('', '').is_available()


def test_implicit_mixed_topic():
    route = interpret_request("Show India's GDP and AI industry growth")
    assert route['structured'] and route['web'].strip() == 'AI industry growth'
    assert interpret_request('Compare India and UAE GDP between 2015 and 2025')['web'] is None


def test_search_again_stays_research(client, search):
    data = client.post('/api/chat', json={'message': 'India GDP latest', 'research_only': True}).get_json()
    assert data['mode'] == 'web'
    search.assert_called_once()


def test_mixed_preserves_data_when_search_fails(client, search, monkeypatch):
    mock_structured(monkeypatch)
    search.side_effect = WebSearchError('Search unavailable')
    data = client.post('/api/chat', json={'message': "Show India's GDP and find information about AI market growth"}).get_json()
    assert data['excel_available'] and data['web']['status'] == 'error'


def test_mixed_preserves_research_when_data_fails(client, search, monkeypatch):
    monkeypatch.setattr(application, 'execute_plan', lambda plan: ([], []))
    data = client.post('/api/chat', json={'message': "Show India's GDP and find information about AI market growth"}).get_json()
    assert not data['excel_available'] and data['web']['results'] and data['structured_error']


def test_limits_and_bad_provider_payload():
    transport = Mock()
    transport.post.return_value.json.return_value = {'results': [{'url': 'https://imf.org/' + str(i)} for i in range(30)]}
    connector = WebSearchConnector('tavily', 'secret', transport)
    assert len(connector.search('AI', 50)['results']) == 10
    transport.post.return_value.json.return_value = []
    with pytest.raises(WebSearchError): connector.search('AI')
    assert not WebSearchConnector('unsupported', 'secret').is_available()
