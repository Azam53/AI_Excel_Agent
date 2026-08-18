from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

GREEN, DARK, PALE, WHITE = "107C41", "12372A", "E9F5EE", "FFFFFF"

def compact(value, value_type):
    if value is None: return "Unavailable"
    if value_type == "currency":
        for size, suffix in ((1e12, "T"), (1e9, "B"), (1e6, "M")):
            if abs(value) >= size: return f"${value/size:,.2f}{suffix}"
        return f"${value:,.0f}"
    if value_type == "population":
        for size, suffix in ((1e9, "B"), (1e6, "M"), (1e3, "K")):
            if abs(value) >= size: return f"{value/size:,.2f}{suffix}"
        return f"{value:,.0f}"
    if value_type == "percent": return f"{value:.2f}%"
    return f"{value:.1f} years"

def generate_workbook(parsed, data, api_urls, generated_at=None):
    now = generated_at or datetime.now()
    wb = Workbook()
    dashboard = wb.active; dashboard.title = "Dashboard"
    ws = wb.create_sheet("Data"); sources = wb.create_sheet("Sources")
    dashboard.sheet_view.showGridLines = False
    dashboard.merge_cells("A1:H2"); dashboard["A1"] = "AI Data Consultant"
    dashboard["A1"].font = Font(size=24, bold=True, color=WHITE); dashboard["A1"].fill = PatternFill("solid", fgColor=DARK); dashboard["A1"].alignment = Alignment(vertical="center")
    dashboard.merge_cells("A3:H3"); dashboard["A3"] = "Public Data Report"; dashboard["A3"].font = Font(size=14, bold=True, color=GREEN)
    country_names = [c["name"] for c in parsed.countries]
    country_label = country_names[0] if len(country_names) == 1 else " vs ".join(country_names)
    dashboard.merge_cells("A5:H5"); dashboard["A5"] = parsed.indicator_name; dashboard["A5"].font = Font(size=22, bold=True, color=DARK)
    dashboard.merge_cells("A6:H6"); dashboard["A6"] = country_label; dashboard["A6"].font = Font(size=14, bold=True, color=GREEN)
    dashboard.merge_cells("A7:H7"); dashboard["A7"] = f"{parsed.start_year} – {parsed.end_year}"; dashboard["A7"].font = Font(size=12, color="587166")
    dashboard.merge_cells("A9:C9"); dashboard["A9"] = "✓ Real public data"; dashboard["A9"].font = Font(bold=True, color=GREEN); dashboard["A9"].fill = PatternFill("solid", fgColor=PALE)
    dashboard.merge_cells("D9:H9"); dashboard["D9"] = "Source: World Bank Open Data"; dashboard["D9"].font = Font(bold=True, color=DARK); dashboard["D9"].fill = PatternFill("solid", fgColor=PALE)
    dashboard["A12"] = "LATEST AVAILABLE DATA"; dashboard["A12"].font = Font(size=14, bold=True, color=DARK)
    headers = ["Country", "Latest Available", "Year"]
    for col, text in enumerate(headers, 1):
        cell = dashboard.cell(13, col, text); cell.font = Font(bold=True, color=WHITE); cell.fill = PatternFill("solid", fgColor=GREEN)
    latest_years = []
    for row, country in enumerate(parsed.countries, 14):
        available = [(year, value) for year, value in data[country["name"]].items() if value is not None]
        latest = max(available, default=(None, None))
        if latest[0] is not None: latest_years.append(latest[0])
        dashboard.cell(row, 1, country["name"]).font = Font(bold=True, color=DARK)
        dashboard.cell(row, 2, compact(latest[1], parsed.value_type)).font = Font(size=13, bold=True, color=GREEN)
        dashboard.cell(row, 3, latest[0] or "Unavailable")
        for col in range(1, 4):
            dashboard.cell(row, col).fill = PatternFill("solid", fgColor="F7FBF8" if row % 2 == 0 else WHITE)
            dashboard.cell(row, col).border = Border(bottom=Side(style="thin", color="D7E3DC"))
    note_row = 15 + len(parsed.countries)
    latest_world_bank_year = max(latest_years, default=None)
    if latest_world_bank_year is not None and latest_world_bank_year < parsed.end_year:
        dashboard.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=3)
        dashboard.cell(note_row, 1, f"Latest World Bank data available: {latest_world_bank_year}")
        dashboard.cell(note_row, 1).font = Font(italic=True, color="7A5A16"); dashboard.cell(note_row, 1).fill = PatternFill("solid", fgColor="FFF8E6")
        dashboard.merge_cells(start_row=note_row + 1, start_column=1, end_row=note_row + 1, end_column=3)
        dashboard.cell(note_row + 1, 1, "Latest available data may be earlier than the requested end year.")
        dashboard.cell(note_row + 1, 1).font = Font(size=9, italic=True, color="7A5A16")
    dashboard.column_dimensions["A"].width = 24; dashboard.column_dimensions["B"].width = 28; dashboard.column_dimensions["C"].width = 14

    ws.append(["Year"] + [c["name"] for c in parsed.countries])
    for year in range(parsed.start_year, parsed.end_year + 1):
        ws.append([year] + [data[c["name"]].get(year) for c in parsed.countries])
    ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]: cell.font = Font(bold=True, color=WHITE); cell.fill = PatternFill("solid", fgColor=GREEN)
    thin = Side(style="thin", color="D7E3DC")
    for row in ws.iter_rows():
        for cell in row: cell.border = Border(bottom=thin); cell.alignment = Alignment(vertical="center")
    for col in range(2, ws.max_column + 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 24
        for row in range(2, ws.max_row + 1):
            ws.cell(row, col).number_format = "$#,##0.00" if parsed.value_type == "currency" else ("0.00\"%\"" if parsed.value_type == "percent" else "#,##0.00")
    table = Table(displayName="ReportData", ref=ws.dimensions); table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
    ws.add_table(table)
    chart = LineChart()
    chart.title = (f"{parsed.indicator_name} — {country_names[0]} — {parsed.start_year} to {parsed.end_year}" if len(parsed.countries) == 1 else f"{parsed.indicator_name} Comparison — {parsed.start_year} to {parsed.end_year}")
    chart.y_axis.title = parsed.indicator_name; chart.x_axis.title = "Year"; chart.style = 13; chart.height = 8; chart.width = 15
    chart.add_data(Reference(ws, min_col=2, max_col=ws.max_column, min_row=1, max_row=ws.max_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)); dashboard.add_chart(chart, "E12")

    source_row = max(note_row + 3, 25)
    dashboard.merge_cells(start_row=source_row, start_column=1, end_row=source_row, end_column=3)
    dashboard.cell(source_row, 1, "Source: World Bank Open Data"); dashboard.cell(source_row, 1).font = Font(size=10, italic=True, color="587166")
    dashboard.merge_cells(start_row=source_row + 1, start_column=1, end_row=source_row + 1, end_column=3)
    dashboard.cell(source_row + 1, 1, f"Generated {now.strftime('%d %b %Y, %H:%M')}"); dashboard.cell(source_row + 1, 1).font = Font(size=9, color="809188")

    source_rows = [("Data Provider", "World Bank"), ("Dataset", "World Bank Open Data"), ("Indicator", parsed.indicator_code), ("Indicator Name", parsed.indicator_name), ("Countries", ", ".join(c["name"] for c in parsed.countries)), ("Period", f"{parsed.start_year} - {parsed.end_year}"), ("Generated", now.isoformat(timespec="seconds"))]
    sources["A1"] = "Report Sources"; sources["A1"].font = Font(size=20, bold=True, color=WHITE); sources["A1"].fill = PatternFill("solid", fgColor=DARK)
    for row, (label, value) in enumerate(source_rows, 3): sources.cell(row, 1, label).font = Font(bold=True); sources.cell(row, 2, value)
    sources.cell(12, 1, "API Requests Used").font = Font(size=14, bold=True, color=GREEN)
    for row, url in enumerate(api_urls, 13): sources.cell(row, 1, url); sources.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    sources.column_dimensions["A"].width = 24; sources.column_dimensions["B"].width = 70
    output = BytesIO(); wb.save(output); output.seek(0); return output

def generate_multi_workbook(context, results, merged, errors=None, generated_at=None):
    now = generated_at or datetime.now(); errors = errors or []
    wb = Workbook(); dashboard = wb.active; dashboard.title = "Dashboard"
    data_ws = wb.create_sheet("Combined Data"); sources = wb.create_sheet("Sources")
    dashboard.sheet_view.showGridLines = False
    dashboard.merge_cells("A1:J2"); dashboard["A1"] = "AI Data Consultant"; dashboard["A1"].font = Font(size=24, bold=True, color=WHITE); dashboard["A1"].fill = PatternFill("solid", fgColor=DARK)
    dashboard.merge_cells("A3:J3"); dashboard["A3"] = "Multi-source Public Data Report"; dashboard["A3"].font = Font(size=14, bold=True, color=GREEN)
    country_names = [country["name"] for country in context["countries"]]
    metric_names = list(dict.fromkeys(result["metric"] for result in results))
    source_names = list(dict.fromkeys(result["source"] for result in results))
    details = [("Period", f'{context["start_year"]} – {context["end_year"]}'), ("Countries", ", ".join(country_names)), ("Metrics", ", ".join(metric_names)), ("Data sources", f'{len(source_names)} — {", ".join(source_names)}'), ("Status", "✓ Real public data — values are not AI-generated")]
    for row, (label, value) in enumerate(details, 5):
        dashboard.cell(row, 1, label).font = Font(bold=True, color=DARK); dashboard.merge_cells(start_row=row, start_column=2, end_row=row, end_column=9); dashboard.cell(row, 2, value)
        if label == "Status": dashboard.cell(row, 2).font = Font(bold=True, color=GREEN); dashboard.cell(row, 2).fill = PatternFill("solid", fgColor=PALE)
    dashboard["A12"] = "LATEST AVAILABLE VALUES"; dashboard["A12"].font = Font(size=14, bold=True, color=DARK)
    for col, value in enumerate(["Series", "Value", "Year", "Source"], 1):
        cell = dashboard.cell(13, col, value); cell.font = Font(bold=True, color=WHITE); cell.fill = PatternFill("solid", fgColor=GREEN)
    from services.data_merger import series_label
    for row, result in enumerate(results, 14):
        available = [(record["period"], record["value"]) for record in result["records"] if record["value"] is not None]
        year, value = max(available, default=(None, None))
        unit_type = "currency" if result["unit"] == "USD" else ("population" if result["unit"] == "people" else ("percent" if result["unit"] == "%" else "years"))
        display = compact(value, unit_type) if result["unit"] in {"USD", "people", "%", "years"} else (f'{value:,.2f} {result["unit"]}' if value is not None else "Unavailable")
        dashboard.append([series_label(result), display, year or "Unavailable", result["source"]])
    for col, width in {"A":30,"B":22,"C":12,"D":18}.items(): dashboard.column_dimensions[col].width = width

    data_ws.append(merged["columns"])
    for row in merged["rows"]: data_ws.append([row.get(column) for column in merged["columns"]])
    data_ws.freeze_panes = "B2"; data_ws.auto_filter.ref = data_ws.dimensions
    for cell in data_ws[1]: cell.font = Font(bold=True, color=WHITE); cell.fill = PatternFill("solid", fgColor=GREEN)
    for column in range(1, data_ws.max_column + 1): data_ws.column_dimensions[data_ws.cell(1, column).column_letter].width = 24
    table = Table(displayName="CombinedReportData", ref=data_ws.dimensions); table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True); data_ws.add_table(table)

    for index, result in enumerate(results[:4]):
        chart = LineChart(); chart.title = series_label(result); chart.y_axis.title = result["unit"]; chart.x_axis.title = "Year"; chart.style = 13; chart.height = 7; chart.width = 12
        chart.add_data(Reference(data_ws, min_col=index + 2, min_row=1, max_row=data_ws.max_row), titles_from_data=True); chart.set_categories(Reference(data_ws, min_col=1, min_row=2, max_row=data_ws.max_row))
        dashboard.add_chart(chart, f'{"F" if index % 2 == 0 else "N"}{12 + (index // 2) * 15}')

    sources["A1"] = "Sources & Methodology"; sources["A1"].font = Font(size=20, bold=True, color=WHITE); sources["A1"].fill = PatternFill("solid", fgColor=DARK)
    row = 3
    for result in results:
        sources.cell(row, 1, "Metric").font = Font(bold=True); sources.cell(row, 2, series_label(result)); row += 1
        sources.cell(row, 1, "Source").font = Font(bold=True); sources.cell(row, 2, result["source"]); row += 1
        for key in ("indicator", "series", "aggregation", "source_url"):
            if result["metadata"].get(key): sources.cell(row, 1, key.replace("_", " ").title()).font = Font(bold=True); sources.cell(row, 2, result["metadata"][key]); row += 1
        row += 1
    if errors:
        sources.cell(row, 1, "Unavailable datasets").font = Font(size=14, bold=True, color="B45309"); row += 1
        for error in errors: sources.cell(row, 1, error["metric"]); sources.cell(row, 2, error["error"]); row += 1
    sources.column_dimensions["A"].width = 24; sources.column_dimensions["B"].width = 100
    output = BytesIO(); wb.save(output); output.seek(0); return output
